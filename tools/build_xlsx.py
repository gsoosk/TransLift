#!/usr/bin/env python3
"""Build the corpus workbook from the CSVs written by aggregate.py.
Usage: build_xlsx.py [deliver dir with CSVs] [output.xlsx]   (defaults: deliver/ and build/java_tx_corpus.xlsx)"""
import sys as _sys
if any(x in ('-h', '--help') for x in _sys.argv[1:]):
    print(__doc__)
    raise SystemExit(0)

import csv, json, os, sys
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import kitlib
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.formula import ArrayFormula

D = sys.argv[1] if len(sys.argv) > 1 else kitlib.DELIVER
OUTFILE = sys.argv[2] if len(sys.argv) > 2 else os.path.join(kitlib.BUILD, 'java_tx_corpus.xlsx')
os.makedirs(os.path.dirname(os.path.abspath(OUTFILE)), exist_ok=True)
FONT = 'Arial'
F_BASE = Font(name=FONT, size=10)
F_BOLD = Font(name=FONT, size=10, bold=True)
F_HEAD = Font(name=FONT, size=10, bold=True, color='FFFFFF')
F_TITLE = Font(name=FONT, size=14, bold=True)
F_SECTION = Font(name=FONT, size=11, bold=True, color='1F4E79')
F_NOTE = Font(name=FONT, size=9, italic=True, color='595959')
FILL_HEAD = PatternFill('solid', fgColor='1F4E79')
FILL_TOTAL = PatternFill('solid', fgColor='DDEBF7')
FILL_INPUT = PatternFill('solid', fgColor='F2F2F2')
WRAP = Alignment(wrap_text=True, vertical='top')
TOP = Alignment(vertical='top')
THIN = Border(top=Side(style='thin', color='808080'))

def read_csv(name):
    path = os.path.join(D, name)
    if not os.path.exists(path):
        return []
    with open(path, newline='') as fh:
        return list(csv.DictReader(fh))

def conv(v, kind):
    if v is None or v == '':
        return None
    if kind == 'int':
        return int(float(v))
    if kind == 'bool':
        return 1 if v == 'True' else 0      # 1/0 rather than TRUE/FALSE: LibreOffice rewrites booleans as =TRUE() formulas
    return v

def header(ws, cols, row=1, widths=None):
    for j, c in enumerate(cols, 1):
        cell = ws.cell(row=row, column=j, value=c)
        cell.font = F_HEAD; cell.fill = FILL_HEAD
        cell.alignment = Alignment(wrap_text=True, vertical='center')
    if widths:
        for j, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(j)].width = w

# ------------------------------------------------------------------ data
apps = read_csv('corpus_summary.csv')
txs = read_csv('transactions.csv')
ops = read_csv('operations.csv')
stats = json.load(open(os.path.join(D, 'corpus_stats.json')))
excl = read_csv('excluded.csv')
N_APPS = len(apps)
REV = stats.get('review') or {}
RECALL = stats.get('recall') or {}

TX_COLS = list(txs[0].keys())
TX_INT = {'line', 'n_ops', 'n_reads', 'n_writes', 'n_tables', 'ops_in_loops', 'writes_in_loops', 'ops_conditional',
          'ops_unknown_table', 'nested_tx_calls', 'app_events', 'cache_ops', 'nondet_clock', 'nondet_id',
          'nondet_random', 'nondet_timer', 'async_calls', 'max_call_depth', 'methods_visited', 'unresolved_calls',
          'polymorphic_skipped', 'invoked_by_count', 'units_total'}
TX_INT &= set(txs[0].keys()) if txs else TX_INT
TX_BOOL = {'read_only', 'non_public', 'db_locking', 'truncated', 'effective', 'inside_db_tx'}
OP_COLS = list(ops[0].keys()) if ops else ['app', 'tx_id', 'tx_kind', 'seq', 'op']
OPL = {c: get_column_letter(i + 1) for i, c in enumerate(OP_COLS)}
OP_INT = {'seq', 'call_depth'}
OP_BOOL = {'in_loop', 'conditional'}
NT = len(txs) + 1          # last data row in Transactions
NO = len(ops) + 1          # last data row in Operations
TXL = {c: get_column_letter(i + 1) for i, c in enumerate(TX_COLS)}

def R(col):
    """Absolute range of a Transactions column."""
    L = TXL[col]
    return f"Transactions!${L}$2:${L}${NT}"

wb = Workbook()

# ------------------------------------------------------------------ README
ws = wb.active
ws.title = 'README'
ws.column_dimensions['A'].width = 30
ws.column_dimensions['B'].width = 120
r = 1
def put(a, b=None, fa=F_BASE, fb=F_BASE):
    global r
    ca = ws.cell(row=r, column=1, value=a); ca.font = fa; ca.alignment = WRAP
    if b is not None:
        cb = ws.cell(row=r, column=2, value=b); cb.font = fb; cb.alignment = WRAP
    r += 1
def section(t):
    global r
    r += 1
    put(t, fa=F_SECTION)

put(f'Java transaction corpus: {N_APPS} open-source apps, extracted for VIC3 chopping analysis', fa=F_TITLE)
ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=2)
ws.cell(row=1, column=1).alignment = Alignment(vertical='center')
ws.row_dimensions[1].height = 24
_rev = (f' Every transaction was then checked by a reviewer agent against the source ({REV.get("confirmed", 0):,} confirmed, '
        f'{REV.get("corrected", 0):,} corrected, {REV.get("added", 0):,} added, {REV.get("excluded", 0):,} excluded'
        + (f', {REV.get("unreviewed", 0):,} not yet reviewed' if REV.get('unreviewed') else '') + ').') if REV else ''
put(f'Built {stats.get("built", "")} with a static extractor (txscan, tree-sitter based) over {N_APPS} Java applications, each pinned to a commit '
    f'(see Apps).{_rev} Every number on the Summary and Apps sheets is a formula over the Transactions sheet, except the '
    'per-app code statistics (grey cells), which come from the extractor.', fa=F_BASE)
ws.merge_cells(start_row=r - 1, start_column=1, end_row=r - 1, end_column=2)
ws.row_dimensions[r - 1].height = 42

section('Sheets')
put('Summary', 'Corpus-level counts by transaction kind, and the shape of explicit transactions that write (size, loops, effects, locks).')
put('Apps', 'One row per application: repository, pinned commit, persistence stack, code size, and transaction counts by kind.')
put('Transactions', f'{len(txs):,} rows, one per transaction: root method, kind, attributes, read/write tables, and chopping-relevant features.')
put('Operations', f'{len(ops):,} rows, one per database operation reached from a transaction root, with tables, flags and SQL (first 300 characters).')
if excl:
    put('Excluded', f'{len(excl):,} extractor candidates that review found are not transactions, with the reason.')

section('What counts as a transaction')
put('explicit', 'The framework runs the code as one database transaction. Declarative: @Transactional (Spring, javax/jakarta) on the method, '
    'an overridden method, or the class; plus app rules (BenchBase Procedure.run). Programmatic: TransactionTemplate, '
    'PlatformTransactionManager, EntityTransaction, JDBC setAutoCommit(false)/commit, JDBI, Seata, Kill Bill\'s transactional wrapper.')
put('implicit', 'An entry point (web endpoint, scheduled job, message or event listener, batch step, servlet, startup hook) that reaches '
    'database operations with no enclosing transaction, so each statement or called transaction commits on its own. '
    'Subkinds: multi-write (two or more write units; a write inside a loop counts as two), read-then-write (a read and a write on '
    'overlapping tables), single-write, read-only.')
put('adhoc', 'Application-level coordination around database operations: synchronized blocks/methods, java.util.concurrent locks, '
    'Redisson/Redis locks, and @Idempotent-style annotations. These are the hand-rolled "transactions" a chopper must respect too.')

section('How it was extracted')
put('Parsing and call graph', 'All non-test Java sources parsed with tree-sitter-java. Calls are resolved by class-hierarchy analysis: for each '
    'concrete subtype of the receiver, the implementation that subtype would run. Calls with more than 8 possible targets are skipped '
    '(counted in polymorphic_skipped).')
put('Walk', 'From each root the extractor follows calls to database operations (max depth 40, 6,000 methods), keeping track of which entity a generic DAO is working on. Each call site counts once per entity. Callees that start their own '
    'transaction (REQUIRES_NEW, NOT_SUPPORTED, NEVER) are recorded in nested_tx_calls and not entered; @Async callees are recorded in async_calls.')
put('Tables', 'From JPA annotations and orm.xml/hbm.xml mappings, Spring Data repositories (derived and @Query), JPQL/HQL, criteria queries, '
    'native SQL, MyBatis XML and annotations, MyBatis-Plus generics, JdbcTemplate/JDBC strings, and Kill Bill .stg templates. '
    'Generic DAOs get the entity bound by the receiver type (e.g. ProgramRuleStore extends IdentifiableObjectStore<ProgramRule>). '
    'Kill Bill @Audited calls also write the entity\'s history table and audit_log, as its JDBI wrapper does. '
    'Table names are lower-cased with schema and quotes removed.')
if RECALL:
    put('Recall check', f'{RECALL["found"]:,} of {RECALL["annotated"]:,} methods annotated @Transactional (with a body, in non-test code) '
        f'are found as roots ({RECALL["recall_pct"]}%). Misses are listed in recall.json.')
if REV:
    put('Review', 'Each transaction was checked against the source by a reviewer agent following docs/review-rules.md: '
        'root and boundary, every operation (statement, tables, read or write), missed operations such as dirty-checking writes, '
        'and the kind. Corrections are applied to the Transactions and Operations sheets (source = review marks added rows; '
        'review_verdict says what the reviewer decided). Confidence is the reviewer\'s own rating.')
else:
    put('Review', 'Not reviewed yet: the rows are the extractor\'s output.')

section('Limitations')
put('Over-approximation', 'Static reachability counts every branch and every possible subtype, so operation counts are upper bounds of what one execution does.')
_worst = ', '.join(stats.get('worst_unknown_table_apps') or [])
put('Unknown tables', f'{stats["ops_unknown_table_pct"]}% of operations have no table: dynamic table names, generic DAOs reached without a type binding, '
    'SQL assembled at runtime.' + (f' Worst: {_worst}' if _worst else '') + ' (see Apps, column "Ops with unknown table %").')
put('JPA dirty checking', 'Updates made through setters on managed entities and written at flush or commit are not listed as operations. '
    f'Flushes are kept but have no table of their own ({stats["ops_without_own_table"]:,} operations are flushes, advisory locks, DDL or sequence calls).')
put('Unresolved dispatch', f'{stats["explicit_no_db_ops_found"]} explicit transactions show no database operation because their calls go through '
    'reflection, external libraries or overly polymorphic calls. Apache Fineract dispatches command handlers by name at runtime, so its '
    'web entry points appear as implicit multi-write transactions over the command log only.')
put('Spring proxies', 'Spring ignores @Transactional on non-public methods and on self-invocations; such roots are kept and marked '
    '(non_public = 1) rather than dropped.')
put('Implicit transactions', 'Approximate what auto-commit does; one entry point can also call several explicit transactions (explicit_tx_called).')

section('Transactions columns')
TX_DOC = {
    'app': 'Application (see Apps).', 'tx_id': 'Unique id <repo>:<n>; joins to Operations.tx_id.',
    'kind': 'explicit, implicit or adhoc.', 'subkind': 'declarative / programmatic-callback / programmatic-region / app-rule; '
    'for implicit <entry kind>/<shape>; for adhoc the coordination construct.',
    'root': 'Class.method where the transaction starts ($L<line> = lambda).', 'file': 'Source file of the root.', 'line': 'Line of the root.',
    'propagation': 'Declared propagation (REQUIRED if unspecified).', 'read_only': '1 if @Transactional(readOnly = true).',
    'isolation': 'Declared isolation level, if any.', 'non_public': '1 if the annotation sits on a non-public method (ignored by Spring proxies).',
    'api': 'Annotation or programmatic API that opens the transaction.', 'n_ops': 'Database operations reached (distinct call site and entity pairs).',
    'n_reads': 'Read operations.', 'n_writes': 'Write operations.', 'n_tables': 'Distinct tables read or written.',
    'tables_read': 'Tables read (;-separated).', 'tables_written': 'Tables written.', 'ops_in_loops': 'Operations inside a loop or stream.',
    'writes_in_loops': 'Writes inside a loop or stream (unbounded number of statements).', 'ops_conditional': 'Operations under a branch or catch.',
    'ops_unknown_table': 'Operations whose table could not be determined.',
    'db_flags': 'Union of operation flags: for_update, pessimistic_lock, advisory_lock, version_check, flush, sequence, '
                'ddl_or_admin, audited (Kill Bill history/audit writes), unknown_table.',
    'db_locking': '1 if it uses database locks (SELECT ... FOR UPDATE, pessimistic lock mode, advisory lock).',
    'versioned_entities': 'Touched entities with an @Version field (optimistic concurrency).',
    'nested_tx_calls': 'Calls into REQUIRES_NEW / NOT_SUPPORTED / NEVER transactions (separate commits).',
    'external_effects': 'Non-database side effects: http, message, email, search-index, object-store (kind x count).',
    'app_events': 'Spring application events published.', 'cache_ops': 'Cache reads/writes/evictions.',
    'nondet_clock': 'Clock reads (now(), currentTimeMillis, new Date).', 'nondet_id': 'ID generation (UUID.randomUUID, snowflake IdWorker).',
    'nondet_random': 'Random numbers.', 'nondet_timer': 'System.nanoTime (usually timing only).',
    'app_locks': 'Application locks reached (ReentrantLock, Semaphore, Redisson RLock, synchronized).',
    'async_calls': '@Async methods or executor submissions reached.', 'max_call_depth': 'Deepest call level with an operation.',
    'methods_visited': 'Methods walked from the root.', 'unresolved_calls': 'Calls to interface methods with no implementation found.',
    'polymorphic_skipped': 'Calls skipped because they had more than 8 targets.', 'truncated': '1 if the walk hit the 6,000-method limit.',
    'invoked_by_count': 'Entry points that reach this transaction.', 'invoked_by': 'First five of those entry points.',
    'units_total': 'Implicit only: auto-committed operations plus explicit transactions called.',
    'key': 'Stable id of the transaction (survives re-numbering); review files are named by it.',
    'source': 'txscan (found by the extractor) or review (added by a reviewer).',
    'review_verdict': 'confirmed, corrected or added; empty if not reviewed.',
    'timeout': 'Declared timeout, if any.', 'rollback_for': 'Declared rollbackFor / rollbackOn.',
    'no_rollback_for': 'Declared noRollbackFor / dontRollbackOn.', 'tx_manager': 'Named transaction manager, if any.',
    'effective': '0 if the reviewer found that the annotation has no effect; empty if not stated.',
    'inside_db_tx': 'Ad hoc only: 1 if the guarded region runs inside a database transaction.',
    'review_confidence': "Reviewer's confidence: high, medium or low.",
    'review_depth': 'full (every operation checked) or sampled (large transaction, operations spot-checked).',
    'explicit_tx_called': 'Implicit only: explicit transactions this entry point calls.',
    'signature': 'Implicit only: operation fingerprint used to count distinct write patterns.',
}
for c in TX_COLS:
    put(c, TX_DOC.get(c, ''))
section('Operations columns')
OP_DOC = {'app': 'Application.', 'tx_id': 'Transaction id (Transactions.tx_id).', 'tx_kind': 'Kind of that transaction.',
          'seq': 'Order in which the walk reached the operation (breadth-first, not execution order).',
          'op': 'R read, W write, L lock, U unknown.', 'api': 'Persistence API (spring-data, jpa, hibernate, mybatis, jdbc, ...).',
          'call': 'Method called.', 'at': 'Class:line of the call.', 'entity': 'Entity the call works on, if known.',
          'tables_read': 'Tables read.', 'tables_written': 'Tables written.', 'flags': 'Operation flags (see db_flags).',
          'in_loop': '1 if inside a loop or stream.', 'conditional': '1 if under a branch or catch.', 'call_depth': 'Call depth from the root.',
          'sql': 'SQL/JPQL text when known (first 300 characters; ? marks parts built at runtime).',
          'source': 'txscan or review (added by a reviewer).', 'review_note': "Reviewer's note on an added or corrected operation."}
for c in OP_COLS:
    put(c, OP_DOC.get(c, ''))

# ------------------------------------------------------------------ Summary (formulas over Transactions / Apps)
ss = wb.create_sheet('Summary')
ss.column_dimensions['A'].width = 74
ss.column_dimensions['B'].width = 14
ss.column_dimensions['C'].width = 70
row = 1
c = ss.cell(row=1, column=1, value='Corpus summary'); c.font = F_TITLE
ss.cell(row=2, column=1, value='All values are formulas over the Transactions, Operations and Apps sheets.').font = F_NOTE
row = 4
refs = {}
def srow(label, formula, fmt='#,##0', note=None, key=None, bold=False, array=False):
    global row
    a = ss.cell(row=row, column=1, value=label); a.font = F_BOLD if bold else F_BASE
    cell = ss.cell(row=row, column=2)
    if array:
        ref = f"B{row}"
        cell.value = ArrayFormula(ref, formula)
    else:
        cell.value = formula
    cell.font = F_BOLD if bold else F_BASE
    cell.number_format = fmt
    if note:
        n = ss.cell(row=row, column=3, value=note); n.font = F_NOTE
    if key:
        refs[key] = f"$B${row}"
    row += 1
def shead(t):
    global row
    row += 1
    h = ss.cell(row=row, column=1, value=t); h.font = F_SECTION
    row += 1

shead('Corpus')
srow('Applications', f'=COUNTA(Apps!$A$2:$A${N_APPS + 1})')
srow('Lines of Java code (non-test)', f'=SUM(Apps!$G$2:$G${N_APPS + 1})')
srow('Transactions', f"=COUNTA({R('tx_id')})", key='tx', bold=True)
srow('Database operations reached from transactions', f"=COUNTA(Operations!$B$2:$B${NO})", key='ops')

shead('Transactions by kind')
srow('Explicit', f'=COUNTIF({R("kind")},"explicit")', key='exp', bold=True)
srow('   declarative (@Transactional, app rules)',
     f'=COUNTIFS({R("kind")},"explicit",{R("subkind")},"declarative")+COUNTIFS({R("kind")},"explicit",{R("subkind")},"app-rule")')
srow('   programmatic (template, transaction manager, JDBC, JDBI, ...)', f'=COUNTIFS({R("kind")},"explicit",{R("subkind")},"programmatic*")')
srow('   with at least one write', f'=COUNTIFS({R("kind")},"explicit",{R("n_writes")},">0")', key='expw')
srow('   marked readOnly', f'=COUNTIFS({R("kind")},"explicit",{R("read_only")},1)')
srow('   no database operation found', f'=COUNTIFS({R("kind")},"explicit",{R("n_ops")},0)',
     note='Calls go through reflection, external libraries or >8-way polymorphic calls.')
srow('   annotation on a non-public method', f'=COUNTIFS({R("kind")},"explicit",{R("non_public")},1)',
     note='Spring proxies ignore these.')
srow('Implicit (entry points without a transaction)', f'=COUNTIF({R("kind")},"implicit")', bold=True)
for sub, note in (('multi-write', 'Two or more write units: each commits separately today.'),
                  ('read-then-write', 'Read and write on overlapping tables without a transaction.'),
                  ('single-write', None), ('read-only', None)):
    srow(f'   {sub}', f'=COUNTIFS({R("kind")},"implicit",{R("subkind")},"*{sub}")', note=note)
srow('Ad hoc (application locks, synchronized, idempotency guards)', f'=COUNTIF({R("kind")},"adhoc")', bold=True)
srow('   synchronized (block or method)', f'=COUNTIFS({R("kind")},"adhoc",{R("subkind")},"synchronized*")')
srow('   application lock (j.u.c., Redisson)', f'=COUNTIFS({R("kind")},"adhoc",{R("subkind")},"app-lock*")')
srow('   annotation (@Idempotent and similar)', f'=COUNTIFS({R("kind")},"adhoc",{R("subkind")},"annotation*")')

shead('Shape of explicit transactions that write')
cond = f'({R("kind")}="explicit")*({R("n_writes")}>0)'
srow('Median database operations', f'=MEDIAN(IF({cond},{R("n_ops")}))', fmt='0', array=True)
srow('90th percentile of database operations', f'=PERCENTILE(IF({cond},{R("n_ops")}),0.9)', fmt='0.0', array=True)
srow('Median tables touched', f'=MEDIAN(IF({cond},{R("n_tables")}))', fmt='0', array=True)
srow('90th percentile of tables touched', f'=PERCENTILE(IF({cond},{R("n_tables")}),0.9)', fmt='0.0', array=True)
def share(label, crit, note=None):
    srow(label, f'=COUNTIFS({R("kind")},"explicit",{R("n_writes")},">0",{crit})/{refs["expw"]}', fmt='0.0%', note=note)
share('Touch a single table', f'{R("n_tables")},"<=1"')
share('Write inside a loop or stream', f'{R("writes_in_loops")},">0"', note='Unbounded number of statements per execution.')
share('Have conditional database operations', f'{R("ops_conditional")},">0"')
share('Call a REQUIRES_NEW / NOT_SUPPORTED transaction', f'{R("nested_tx_calls")},">0"', note='Already split into separate commits.')
share('Have external side effects (HTTP, messaging, email, search, object store)', f'{R("external_effects")},"<>"')
share('Publish in-process application events', f'{R("app_events")},">0"')
share('Use a cache', f'{R("cache_ops")},">0"')
srow('Read the clock, generate IDs or random numbers',
     f'=SUMPRODUCT({cond}*(({R("nondet_clock")}+{R("nondet_id")}+{R("nondet_random")})>0))/{refs["expw"]}', fmt='0.0%')
share('Hand work to @Async methods or executors', f'{R("async_calls")},">0"')
share('Use database locks (FOR UPDATE, pessimistic, advisory)', f'{R("db_locking")},1')
share('Hold application locks', f'{R("app_locks")},"<>"')
share('Touch an @Version (optimistically locked) entity', f'{R("versioned_entities")},"<>"')
share('Have at least one operation on an unknown table', f'{R("ops_unknown_table")},">0"')

shead('Operations')
srow('Operations with an unknown table', f'=SUM({R("ops_unknown_table")})', key='unk')
srow('Share of operations with an unknown table', f'={refs["unk"]}/{refs["ops"]}', fmt='0.0%')
srow('Operations inside loops', f'=COUNTIF(Operations!${OPL["in_loop"]}$2:${OPL["in_loop"]}${NO},1)')
srow('Write operations', f'=COUNTIF(Operations!${OPL["op"]}$2:${OPL["op"]}${NO},"W")')
if 'review_verdict' in TXL:
    shead('Review')
    srow('Confirmed as extracted', f'=COUNTIF({R("review_verdict")},"confirmed")')
    srow('Corrected by the reviewer', f'=COUNTIF({R("review_verdict")},"corrected")')
    srow('Added by the reviewer (missed by the extractor)', f'=COUNTIF({R("review_verdict")},"added")')
    srow('Not reviewed', f'=COUNTBLANK({R("review_verdict")})')
    srow('Excluded by the reviewer (not transactions; see Excluded)', len(excl))
    srow('Reviewed with low confidence', f'=COUNTIF({R("review_confidence")},"low")')

# ------------------------------------------------------------------ Apps
wa = wb.create_sheet('Apps')
A_COLS = ['App', 'Repository', 'Commit', 'Domain', 'Persistence', 'Java files', 'Lines of code', 'Entities',
          'MyBatis statements', 'DB call sites', 'Entry points', 'Explicit: declarative', 'Explicit: programmatic',
          'Implicit: multi-write', 'Implicit: read-then-write', 'Implicit: single-write', 'Implicit: read-only', 'Ad hoc',
          'Transactions', 'Explicit with writes', 'Median ops per explicit write tx', 'Explicit with writes in loops',
          'Operations', 'Ops with unknown table', 'Ops with unknown table %', '@Version entities', 'App lock sites',
          'synchronized methods', 'Stars', 'Selected by', 'Review: confirmed', 'Review: corrected',
          'Review: added', 'Review: excluded', 'Review: not reviewed']
header(wa, A_COLS, widths=[20, 34, 13, 30, 28, 9, 11, 9, 10, 10, 10, 11, 12, 11, 12, 11, 11, 8, 12, 11, 12, 12, 11, 11, 11, 10, 9, 12,
                           9, 10, 11, 11, 10, 10, 11])
wa.row_dimensions[1].height = 45
STATIC = [('app', None), ('repo', None), ('commit', None), ('domain', None), ('persistence', None), ('java_files', 'int'),
          ('lines_of_code', 'int'), ('entities', 'int'), ('mybatis_statements', 'int'), ('db_call_sites', 'int'),
          ('entry_points', 'int')]
for i, a in enumerate(apps, 2):
    for j, (k, kind) in enumerate(STATIC, 1):
        cell = wa.cell(row=i, column=j, value=conv(a[k], kind))
        cell.font = F_BASE
        if j >= 6:
            cell.fill = FILL_INPUT; cell.number_format = '#,##0'
    A = f'$A{i}'
    f = {
        12: f'=COUNTIFS({R("app")},{A},{R("kind")},"explicit",{R("subkind")},"declarative")+COUNTIFS({R("app")},{A},{R("kind")},"explicit",{R("subkind")},"app-rule")',
        13: f'=COUNTIFS({R("app")},{A},{R("kind")},"explicit",{R("subkind")},"programmatic*")',
        14: f'=COUNTIFS({R("app")},{A},{R("kind")},"implicit",{R("subkind")},"*multi-write")',
        15: f'=COUNTIFS({R("app")},{A},{R("kind")},"implicit",{R("subkind")},"*read-then-write")',
        16: f'=COUNTIFS({R("app")},{A},{R("kind")},"implicit",{R("subkind")},"*single-write")',
        17: f'=COUNTIFS({R("app")},{A},{R("kind")},"implicit",{R("subkind")},"*read-only")',
        18: f'=COUNTIFS({R("app")},{A},{R("kind")},"adhoc")',
        19: f'=SUM(L{i}:R{i})',
        20: f'=COUNTIFS({R("app")},{A},{R("kind")},"explicit",{R("n_writes")},">0")',
        22: f'=COUNTIFS({R("app")},{A},{R("kind")},"explicit",{R("n_writes")},">0",{R("writes_in_loops")},">0")',
        23: f'=SUMIFS({R("n_ops")},{R("app")},{A})',
        24: f'=SUMIFS({R("ops_unknown_table")},{R("app")},{A})',
        25: f'=IF(W{i}=0,0,X{i}/W{i})',
    }
    for col, formula in f.items():
        cell = wa.cell(row=i, column=col, value=formula); cell.font = F_BASE
        cell.number_format = '0.0%' if col == 25 else '#,##0'
    cell = wa.cell(row=i, column=21)
    cell.value = ArrayFormula(f'U{i}', f'=IF(T{i}=0,0,MEDIAN(IF(({R("app")}={A})*({R("kind")}="explicit")*({R("n_writes")}>0),{R("n_ops")})))')
    cell.font = F_BASE; cell.number_format = '0'
    for col, k in ((26, 'versioned_entities'), (27, 'app_lock_sites'), (28, 'sync_methods'), (29, 'stars'),
                   (31, 'review_confirmed'), (32, 'review_corrected'), (33, 'review_added'), (34, 'review_excluded'),
                   (35, 'review_unreviewed')):
        cell = wa.cell(row=i, column=col, value=conv(a.get(k, ''), 'int')); cell.font = F_BASE; cell.fill = FILL_INPUT
        cell.number_format = '#,##0'
    cell = wa.cell(row=i, column=30, value=a.get('source', '')); cell.font = F_BASE; cell.fill = FILL_INPUT
last = len(apps) + 1
tr = last + 1
wa.cell(row=tr, column=1, value='All apps').font = F_BOLD
for col in list(range(6, 21)) + [22, 23, 24, 26, 27, 28, 31, 32, 33, 34, 35]:
    L = get_column_letter(col)
    cell = wa.cell(row=tr, column=col, value=f'=SUM({L}2:{L}{last})'); cell.number_format = '#,##0'
cell = wa.cell(row=tr, column=21)
cell.value = ArrayFormula(f'U{tr}', f'=MEDIAN(IF(({R("kind")}="explicit")*({R("n_writes")}>0),{R("n_ops")}))')
cell.number_format = '0'
wa.cell(row=tr, column=25, value=f'=IF(W{tr}=0,0,X{tr}/W{tr})').number_format = '0.0%'
for col in range(1, len(A_COLS) + 1):
    cell = wa.cell(row=tr, column=col); cell.font = F_BOLD; cell.fill = FILL_TOTAL; cell.border = THIN
wa.cell(row=tr + 2, column=1, value='Grey cells come from the extractor (code statistics); all other numbers are formulas over the Transactions sheet.').font = F_NOTE
wa.freeze_panes = 'B2'

# ------------------------------------------------------------------ Transactions
wt = wb.create_sheet('Transactions')
TW = {'app': 18, 'tx_id': 30, 'kind': 9, 'subkind': 22, 'root': 46, 'file': 50, 'tables_read': 40, 'tables_written': 40,
      'db_flags': 16, 'versioned_entities': 24, 'external_effects': 20, 'app_locks': 24, 'invoked_by': 50,
      'explicit_tx_called': 30, 'signature': 50, 'api': 18}
header(wt, TX_COLS, widths=[TW.get(c, 10) for c in TX_COLS])
wt.row_dimensions[1].height = 30
for i, t in enumerate(txs, 2):
    for j, c in enumerate(TX_COLS, 1):
        kind = 'int' if c in TX_INT else ('bool' if c in TX_BOOL else None)
        wt.cell(row=i, column=j, value=conv(t[c], kind))
wt.freeze_panes = 'C2'
wt.auto_filter.ref = f"A1:{get_column_letter(len(TX_COLS))}{NT}"

# ------------------------------------------------------------------ Operations
wo = wb.create_sheet('Operations')
OW = {'app': 18, 'tx_id': 30, 'tx_kind': 9, 'op': 5, 'api': 13, 'call': 26, 'at': 34, 'entity': 26, 'tables_read': 30,
      'tables_written': 30, 'flags': 16, 'sql': 90}
header(wo, OP_COLS, widths=[OW.get(c, 9) for c in OP_COLS])
for i, o in enumerate(ops, 2):
    for j, c in enumerate(OP_COLS, 1):
        kind = 'int' if c in OP_INT else ('bool' if c in OP_BOOL else None)
        wo.cell(row=i, column=j, value=conv(o[c], kind))
wo.freeze_panes = 'C2'
wo.auto_filter.ref = f"A1:{get_column_letter(len(OP_COLS))}{NO}"

if excl:
    wx = wb.create_sheet('Excluded')
    XC = list(excl[0].keys())
    header(wx, XC, widths=[18, 30, 14, 9, 22, 46, 50, 7, 90])
    for i, x in enumerate(excl, 2):
        for j, c in enumerate(XC, 1):
            wx.cell(row=i, column=j, value=conv(x[c], 'int' if c == 'line' else None))
    wx.freeze_panes = 'C2'

wb.save(OUTFILE)
print('saved', OUTFILE, 'tx rows', len(txs), 'op rows', len(ops))
