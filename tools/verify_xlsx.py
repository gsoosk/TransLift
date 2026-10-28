#!/usr/bin/env python3
"""Check the workbook's formula results against aggregate.py's statistics (after tools/recalc_xlsx.sh).
Usage: verify_xlsx.py [deliver dir] [xlsx]   (defaults: deliver/ and build/java_tx_corpus.xlsx)"""
import sys as _sys
if any(x in ('-h', '--help') for x in _sys.argv[1:]):
    print(__doc__)
    raise SystemExit(0)

import json, csv, os, sys
from openpyxl import load_workbook
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import kitlib
D = sys.argv[1] if len(sys.argv) > 1 else kitlib.DELIVER
X = sys.argv[2] if len(sys.argv) > 2 else os.path.join(kitlib.BUILD, 'java_tx_corpus.xlsx')
wb = load_workbook(X, data_only=True)
vals = {r[0].strip(): r[1] for r in wb['Summary'].iter_rows(min_row=4, values_only=True) if r[0] is not None and r[1] is not None}
if vals.get('Transactions') is None:          # formulas were never calculated
    print('the workbook has no computed values: run tools/recalc_xlsx.sh first (needs LibreOffice)')
    sys.exit(1)
st = json.load(open(os.path.join(D, 'corpus_stats.json')))
checks = [('Transactions', 'transactions_total', 1), ('Explicit', 'explicit', 1), ('Implicit (entry points without a transaction)', 'implicit', 1),
          ('Ad hoc (application locks, synchronized, idempotency guards)', 'adhoc', 1), ('Database operations reached from transactions', 'ops_total', 1),
          ('Median database operations', 'explicitW_ops_median', 1), ('90th percentile of database operations', 'explicitW_ops_p90', 1),
          ('Median tables touched', 'explicitW_tables_median', 1), ('90th percentile of tables touched', 'explicitW_tables_p90', 1),
          ('Touch a single table', 'explicitW_single_table_pct', 100), ('Write inside a loop or stream', 'explicitW_with_loop_writes_pct', 100),
          ('Have conditional database operations', 'explicitW_with_conditional_ops_pct', 100),
          ('Call a REQUIRES_NEW / NOT_SUPPORTED transaction', 'explicitW_with_requires_new_pct', 100),
          ('Have external side effects (HTTP, messaging, email, search, object store)', 'explicitW_with_external_effects_pct', 100),
          ('Publish in-process application events', 'explicitW_with_app_events_pct', 100), ('Use a cache', 'explicitW_with_cache_ops_pct', 100),
          ('Read the clock, generate IDs or random numbers', 'explicitW_with_clock_id_random_pct', 100),
          ('Hand work to @Async methods or executors', 'explicitW_with_async_pct', 100),
          ('Use database locks (FOR UPDATE, pessimistic, advisory)', 'explicitW_with_db_locking_pct', 100),
          ('Hold application locks', 'explicitW_with_app_locks_pct', 100),
          ('Touch an @Version (optimistically locked) entity', 'explicitW_touching_versioned_entity_pct', 100),
          ('Have at least one operation on an unknown table', 'explicitW_with_unknown_table_ops_pct', 100),
          ('Share of operations with an unknown table', 'ops_unknown_table_pct', 100)]
bad = 0
for label, key, mult in checks:
    v = vals.get(label)
    ok = v is not None and abs(round(v * mult, 1) - round(st[key], 1)) < 0.051
    bad += not ok
    if not ok:
        print('BAD', label, v, st[key])
wa = wb['Apps']
hdr = [c.value for c in wa[1]]
apps = list(csv.DictReader(open(os.path.join(D, 'corpus_summary.csv'))))
pairs = {'Explicit: declarative': 'explicit_declarative', 'Explicit: programmatic': 'explicit_programmatic',
         'Implicit: multi-write': 'implicit_multi_write', 'Implicit: read-then-write': 'implicit_read_then_write',
         'Implicit: single-write': 'implicit_single_write', 'Implicit: read-only': 'implicit_read_only', 'Ad hoc': 'adhoc',
         'Explicit with writes': 'explicit_with_writes', 'Operations': 'ops_total', 'Ops with unknown table': 'ops_unknown_table'}
mism = sum(1 for i, a in enumerate(apps, 2) for h, k in pairs.items() if int(wa.cell(row=i, column=hdr.index(h) + 1).value) != int(a[k]))
print(f'summary checks failed: {bad} of {len(checks)}; app cell mismatches: {mism}')
sys.exit(1 if bad or mism else 0)
